#!/usr/bin/env python3
"""Build architecture-plan.pptx — every HTML block as a 16:9 slide (screenshot, no summarization)."""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

DOCS = Path(__file__).resolve().parent
HTML = DOCS / "architecture-plan.html"
CSS = DOCS / "export-slide.css"
BUILDER_JS = DOCS / "export-slide-builder.js"
PPTX = DOCS / "architecture-plan.pptx"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
VIEWPORT_W = 1280
VIEWPORT_H = 720


def ensure_deps() -> None:
    try:
        import pptx  # noqa: F401
    except ImportError:
        subprocess = __import__("subprocess")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])


def build_pptx() -> int:
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from pptx.util import Inches

    css = CSS.read_text(encoding="utf-8")
    js = BUILDER_JS.read_text(encoding="utf-8")

    work = Path(tempfile.mkdtemp(prefix="arch-ppt-"))
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": VIEWPORT_W, "height": VIEWPORT_H})
            page.goto(HTML.as_uri(), wait_until="networkidle")
            page.add_style_tag(content=css)
            page.add_script_tag(content=js)
            page.wait_for_timeout(900)

            slides = page.locator(".ppt-export-slide")
            count = slides.count()
            if count == 0:
                raise RuntimeError("export-slide-builder.js produced no slides")

            prs = Presentation()
            prs.slide_width = Inches(SLIDE_W_IN)
            prs.slide_height = Inches(SLIDE_H_IN)
            blank = prs.slide_layouts[6]

            for i in range(count):
                png = work / f"slide-{i:03d}.png"
                slides.nth(i).screenshot(path=str(png))
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(
                    str(png),
                    Inches(0),
                    Inches(0),
                    width=Inches(SLIDE_W_IN),
                    height=Inches(SLIDE_H_IN),
                )

            browser.close()

        prs.save(str(PPTX))
        return count
    finally:
        shutil.rmtree(work, ignore_errors=True)


def main() -> None:
    for path, label in (
        (HTML, "architecture-plan.html"),
        (CSS, "export-slide.css"),
        (BUILDER_JS, "export-slide-builder.js"),
    ):
        if not path.is_file():
            raise SystemExit(f"Missing {label}")

    ensure_deps()
    count = build_pptx()
    print(f"Wrote {PPTX} ({count} slides, full HTML content · 16:9)")


if __name__ == "__main__":
    main()
